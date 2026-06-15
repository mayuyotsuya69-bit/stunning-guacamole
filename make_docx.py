#!/usr/bin/env python3
"""AI関与スコア 説明資料の Word(.docx) 版を生成する（報告書レイアウト）。

トンマナ: 白背景 / 黒本文 / 見出しアクセント #3E8C7C /
          ヒラギノ角ゴ ProN（本文 W3・見出し W6）。
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x3E, 0x8C, 0x7C)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
W3 = "ヒラギノ角ゴ ProN W3"
W6 = "ヒラギノ角ゴ ProN W6"

doc = Document()

# 既定スタイル
normal = doc.styles["Normal"]
normal.font.name = W3
normal.font.size = Pt(10.5)
normal.font.color.rgb = BLACK
normal.element.rPr.rFonts.set(qn("w:eastAsia"), W3)


def _ea(run, name):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts"); rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


def para(text="", size=10.5, color=BLACK, font=W3, before=0, after=4, align=None, bold=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    pf.line_spacing = 1.15
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size); r.font.color.rgb = color; r.bold = bold
        _ea(r, font)
    return p


def heading(text, num=None, size=15):
    label = f"{num}. {text}" if num else text
    p = para(label, size=size, color=ACCENT, font=W6, before=12, after=4)
    # 下罫線（アクセント）
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4"); bottom.set(qn("w:color"), "3E8C7C")
    pbdr.append(bottom); pPr.append(pbdr)
    return p


def bullet(text, color=BLACK, font=W3, size=10.5, sub=False):
    p = para(after=3)
    pf = p.paragraph_format
    pf.left_indent = Cm(1.4 if sub else 0.8)
    pf.first_line_indent = Cm(-0.5)
    r = p.add_run(("－ " if sub else "● ") + text)
    r.font.size = Pt(size); r.font.color.rgb = color; _ea(r, font)
    return p


# ===== タイトル =====
para("AI関与スコア　算出方法 説明資料", size=22, color=BLACK, font=W6, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
para("作文が生成AIで書かれた可能性を、どのように判定しているか", size=12, color=ACCENT, font=W6, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
para("社長決裁資料　／　2026年6月15日　／　位置づけ：参考指標（断定ではない）",
     size=10, color=BLACK, font=W3, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)

# ===== 0. 結論 =====
heading("結論（先に要点）", "0")
bullet("「AIが書いたか／自分で考えて書いたか」を最終的に断定するものではありません。", font=W6)
bullet("生成AIへの「丸写し・大部分の写し」の“可能性の高さ”を 0〜100 で示す参考指標です。", sub=True)
bullet("数値は統計式や検出器による機械的な計算ではありません。", font=W6)
bullet("判定用の生成AI（LLM）が、全作文を“同一の評価基準”で読み、点数を付けています。", sub=True)
bullet("＝「採点ルールを固定した採点者（AI）に、全員分を同じ物差しで読ませている」仕組みです。", sub=True)
bullet("判定の核心は、次の3点を見ていることです。", font=W6)
bullet("本人の体験・実感・固有性が見えるか", sub=True)
bullet("文章が整いすぎていないか", sub=True)
bullet("AI特有の抽象論・定型表現・均質な構成が続いていないか", sub=True)

# ===== 1. 目的 =====
heading("目的", "1")
bullet("手書き提出の作文について、生成AIの文章を「そのまま／大部分写している可能性」を見立てる指標です。")
bullet("手書きであること自体は、人間が書いた根拠として扱いません（手書きでもAI出力を写せるため）。")
bullet("判定対象は、文章の「内容・構成・語彙・表現」のみです。")
bullet("目的はスクリーニング（要確認の作文を見つけること）であり、最終判断は人が行う前提です。", font=W6)

# ===== 2. スコアの意味 =====
heading("スコアの意味（0〜100の整数）", "2")
tbl = doc.add_table(rows=5, cols=2)
tbl.style = "Table Grid"
tbl.columns[0].width = Cm(3.5); tbl.columns[1].width = Cm(12.5)
rows = [("スコア帯", "意味"),
        ("0〜30", "書き写した可能性は低い"),
        ("31〜60", "一部参考にした／判定困難"),
        ("61〜80", "書き写した可能性がやや高い"),
        ("81〜100", "書き写した可能性が高い")]
for i, (a, b) in enumerate(rows):
    for j, txt in enumerate((a, b)):
        cell = tbl.cell(i, j)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        r = cell.paragraphs[0].add_run(txt)
        r.font.size = Pt(10.5)
        r.font.color.rgb = (RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else BLACK)
        r.bold = (i == 0)
        _ea(r, W6 if i == 0 else W3)
        if i == 0:
            shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "3E8C7C")
            cell._tc.get_or_add_tcPr().append(shd)
para("運用方針：迷う場合は 40〜60 に寄せます。0〜10 や 90〜100 の極端な値は、明確な根拠が複数そろう場合のみ使用します。",
     size=10, color=BLACK, before=4)

# ===== 3. 判定の仕組み =====
heading("判定の仕組み（どう点数を出しているか）", "3")
para("作文（手書き） → テキスト化 → 判定用AI（LLM）＋評価基準 → スコア 0〜100 ＋根拠3点",
     size=11, color=ACCENT, font=W6, after=4)
bullet("スコアは採点者役のAIの“判断”であり、固定の計算式から自動算出される値ではありません。")
bullet("判断がブレないよう、評価の観点・各スコア帯の定義・禁止事項を文章で細かく固定しています。")
bullet("→「全員を同じ物差しで読む」ことを担保しています。", color=ACCENT, font=W6, sub=True)
bullet("同じ作文でも判定は多少揺れ得るため、確定値ではなく“目安”として扱います。", font=W6)

# ===== 4. 評価の3軸 =====
heading("評価の3つの軸（AIは何を見ているか）", "4")
axes = [
    ("A. 具体性・本人性",
     "経験・失敗・迷いの記述／部活・家族・友人・地域／固有名詞・具体的な場面",
     "一般論・大きな主語ばかり／どの生徒でも書ける内容／体験が抽象的で場面が不明"),
    ("B. 文体・語彙",
     "自然な言い回し／文の長さにばらつき／感情や言い直しの揺れ",
     "均質すぎる文体／「〜が重要である」等の定型表現／不自然に硬い抽象語の多用"),
    ("C. 構成",
     "多少の偏り・寄り道がある／段落ごとに熱量差／本人の考えの流れが見える",
     "整いすぎた序論・本論・結論／機械的な整理／模範解答的すぎる結び"),
]
for name, low, high in axes:
    bullet(name, font=W6)
    bullet("人間らしい（可能性は低）：" + low, color=RGBColor(0x2E, 0x7D, 0x57), sub=True)
    bullet("AIらしい（可能性は高）：" + high, color=RGBColor(0xB3, 0x3A, 0x2E), sub=True)

# ===== 5. 誤判定を防ぐルール =====
heading("誤判定を防ぐルール（公平性の担保）", "5")
bullet("単独の特徴だけでは高スコアにしません。", font=W6)
for t in ["「手書きだから人間」とは判断しない",
          "「誤字脱字があるから人間」とは単独で判断しない",
          "「文章が上手いからAI」とは判断しない",
          "「構成が整っているからAI」とは単独で判断しない",
          "1つの特徴だけで極端なスコアにしない／「AIが書いた」と断定しない"]:
    bullet(t, sub=True)
bullet("高スコアにしてよいのは、複数の根拠がそろう場合のみです。", color=ACCENT, font=W6)

# ===== 6. 出力形式 =====
heading("出力形式（1件ごと）", "6")
bullet("スコアだけでなく根拠3点をセットで残し、後から人が確認・反証できるようにしています。")
para("スコア: <0〜100の整数>\n根拠:\n  - <根拠1>\n  - <根拠2>\n  - <根拠3>",
     size=10.5, color=BLACK, font="Osaka-Mono", before=2, after=4)

# ===== 7. 限界 =====
heading("この指標の限界（必ず共有すべき点）", "7")
bullet("断定ではなく“可能性”の参考値です。最終判断は必ず人が行います。", font=W6)
bullet("AI判定のため、同じ作文でもスコアが多少揺れることがあります。")
bullet("文章力が高い生徒を過剰に疑わない設計ですが、誤判定を完全には排除できません。")
bullet("高スコアの作文は「クロ」ではなく、「人による確認を優先すべき作文」と位置づけます。", color=ACCENT, font=W6)

# ===== 8. 想定問答 =====
heading("想定問答（社長質問への回答案）", "8")
qa = [
    ("Q. どうやってAIが書いたか／自分で考えたかを判断しているのか？",
     "A. 専用の生成AIに全作文を同一基準で読ませ、「本人の体験が見えるか」「整いすぎていないか」「AI特有の定型表現や均質な構成が続いていないか」の3観点から、0〜100のスコアと根拠3点を出力します。機械的な計算ではなく、基準を固定した採点者（AI）による見立てです。"),
    ("Q. 文章がうまい生徒を、AI扱いしてしまわないか？",
     "A. 「上手い・整っている・誤字が少ない」といった単独の特徴だけでは高スコアにしないルールを明示しています。高スコアになるのは、複数の根拠がそろった場合のみです。"),
    ("Q. このスコアで処分や合否を決めてよいか？",
     "A. いいえ。スクリーニング用の参考指標です。高スコアの作文を抽出し、最終的には根拠と原文を人が確認して判断する運用を想定しています。"),
]
for q, a in qa:
    para(q, size=11, color=ACCENT, font=W6, before=6, after=2)
    para(a, size=10.5, color=BLACK, font=W3, after=2)

doc.save("AI関与スコア_算出方法_説明資料.docx")
print("docx saved")
